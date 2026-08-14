# -*- coding: utf-8 -*-
"""
Reprocessamento retroativo das rodadas anteriores (seção 12).
=============================================================

Roda o motor corrigido sobre as rodadas já emitidas e compara o resultado com o
Termo que de fato saiu. É o que transforma "sabemos de R$ 10.367,84 na 11ª e
R$ 9.789,13 na 12ª" em uma lista nominal fechada, título a título — e é ela que
encerra a divergência de contagem 5 × 6 × 7 do adendo 17.

**O OneDrive é somente leitura** (regra 15.8). Cada arquivo é copiado para
``snapshots/onedrive/`` e o processamento acontece sobre a cópia. Nada é
escrito, movido ou renomeado na pasta de origem.

O que o layout dos arquivos permite
-----------------------------------

A estrutura da planilha mudou várias vezes ao longo de 2026, e isso limita o
que dá para apurar em cada rodada:

* **completa** — abas ``VORTX`` e ``GRAFENO`` presentes. Dá para reconciliar do
  zero e comparar com o Termo emitido. São as rodadas da 9ª em diante;
* **só o Termo** — há ``TERMO DE RECOMPRA`` mas não as duas bases. Dá para
  conferir duplicidade dentro do Termo, e só;
* **sem estrutura** — nem Termo. Fica registrada como não processada, com o
  motivo. Não se inventa número onde não há dado.

A ordem importa
---------------

As rodadas são processadas em ordem crescente e o histórico vai sendo alimentado
com o Termo **efetivamente emitido** de cada uma. Assim, ao chegar na rodada
seguinte, a trava contra recompra repetida enxerga o que já havia sido cobrado —
que é exatamente como o título 165955407 (Cíntia Farias Cordeiro) foi cobrado
três vezes entre a 11ª e a 12ª.
"""

import re
import shutil
import tempfile
from decimal import Decimal
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font

from core import calculo, config, leitura, motor
from core import normalizacao as nz
from core.excecoes import _DATA, _linha, _nova_aba

PASTA_COPIAS = config.PASTA_SNAPSHOTS / "onedrive"

#  Arquivos que não são a versão oficial da rodada.
_DESCARTAR = ("copia", "cópia", "modelo", "backup", "antiga", "teste")

SITUACAO_COMPLETA = "completa"
SITUACAO_SO_TERMO = "só o Termo"
SITUACAO_SEM_ESTRUTURA = "sem estrutura"


# ---------------------------------------------------------------------------
#  Levantamento das rodadas
# ---------------------------------------------------------------------------
class Rodada:
    """Uma pasta de rodada do OneDrive e o que dá para fazer com ela."""

    def __init__(self, numero, pasta, arquivo, abas):
        self.numero = numero
        self.pasta = Path(pasta)
        self.arquivo = Path(arquivo) if arquivo else None
        self.abas = abas or []
        self.situacao = self._classificar()
        self.motivo = ""
        #  preenchidos pelo reprocessamento
        self.copia = None
        self.parametros = None
        self.emitidos = []          # dicionários, na ordem do Termo emitido
        self.id_confiavel = False   # a coluna de número identifica o título?
        self.criterio_duplicidade = ""
        self.titulos_distintos = 0
        self.vezes_por_chave = {}
        self.calculados = {}        # id -> linha do motor
        self.faltantes = []
        self.indevidos = []
        self.duplicados = []
        self.contagens = {}
        #  Motivo já resolvido em texto por título. O objeto ``Resultado`` inteiro
        #  não fica guardado: são 62 mil títulos e 142 mil boletos por rodada, e
        #  segurar quatro rodadas ao mesmo tempo estouraria a memória. O que o
        #  relatório precisa é a frase, não a base.
        self.causas = {}

    def _classificar(self):
        nomes = {_sem_acento(a) for a in self.abas}
        if {"VORTX", "GRAFENO"} <= nomes and "TERMO DE RECOMPRA" in nomes:
            return SITUACAO_COMPLETA
        if "TERMO DE RECOMPRA" in nomes:
            return SITUACAO_SO_TERMO
        return SITUACAO_SEM_ESTRUTURA

    @property
    def valor_faltante(self):
        return calculo.somar(l["valor_recompra"] for l in self.faltantes)

    @property
    def valor_indevido(self):
        return calculo.somar(l["valor_recompra"] for l in self.indevidos)

    @property
    def valor_duplicado(self):
        return calculo.somar(l["valor_recompra"] for l in self.duplicados)

    def __repr__(self):
        return (f"<Rodada {self.numero} {self.situacao} "
                f"{self.arquivo.name if self.arquivo else '—'}>")


def _sem_acento(texto):
    return nz.sem_acento(str(texto or "")).upper().strip()


def _versao(caminho):
    """Ordena as versões do mesmo arquivo: ``v5.1`` vence ``v4.1`` e ``v1.1``."""
    achado = re.search(r"_v(\d+)[._](\d+)", caminho.name)
    if achado:
        return (int(achado.group(1)), int(achado.group(2)))
    achado = re.search(r"_v(\d+)", caminho.name)
    return (int(achado.group(1)), 0) if achado else (0, 0)


def _numero_da_pasta(pasta):
    """
    O número da rodada, deduzido do nome da pasta.

    As pastas são numeradas a partir de ``02.Recompra_15_01_26``, que é a 1ª
    rodada — ``01.Outros`` não é rodada. O nome do arquivo confirma onde ele o
    traz (``Leve_Saude_9_Recompra_...`` na pasta ``10.``).
    """
    achado = re.match(r"(\d+)\.", pasta.name)
    if not achado or not pasta.name.lower().count("recompra"):
        return None
    return int(achado.group(1)) - 1


def _escolher_arquivo(pasta):
    """A planilha da rodada: a de maior versão, preferindo o nome oficial."""
    candidatos = [p for p in pasta.glob("*.xls*")
                  if not p.name.startswith("~$")
                  and not any(m in p.name.lower() for m in _DESCARTAR)]
    if not candidatos:
        return None
    oficiais = [p for p in candidatos if p.name.lower().startswith("leve_saude")]
    return max(oficiais or candidatos, key=lambda p: (_versao(p), p.stat().st_size))


def abas_do_arquivo(caminho):
    """
    Nomes das abas, lendo só ``xl/workbook.xml`` de dentro do ``.xlsx``.

    Alguns arquivos passam de 50 MB; abrir com o openpyxl só para listar aba
    custaria minutos por rodada. E é leitura pura — o original nunca é aberto
    para escrita.
    """
    import zipfile
    try:
        with zipfile.ZipFile(caminho) as z:
            wb = z.read("xl/workbook.xml").decode("utf-8", "replace")
    except (zipfile.BadZipFile, KeyError, OSError):
        return []
    return [_desescapar(n) for n in re.findall(r'<sheet[^>]*name="([^"]+)"', wb)]


def _desescapar(s):
    return (s.replace("&amp;", "&").replace("&lt;", "<").replace("&gt;", ">")
            .replace("&quot;", '"').replace("&apos;", "'"))


def varrer(raiz=None, log=None):
    """Lista as rodadas do OneDrive, da mais antiga para a mais nova."""
    log = log or (lambda *a, **k: None)
    raiz = Path(raiz or config.PASTA_ONEDRIVE)
    if not raiz.exists():
        raise FileNotFoundError(
            f"Pasta das rodadas anteriores não encontrada:\n{raiz}\n"
            f"Confira o caminho ou se o OneDrive está sincronizado."
        )

    log(f"Varrendo {raiz}…")
    rodadas = []
    for pasta in sorted(raiz.iterdir()):
        if not pasta.is_dir():
            continue
        numero = _numero_da_pasta(pasta)
        if numero is None:
            log(f"  {pasta.name}: não é pasta de rodada — ignorada.")
            continue
        arquivo = _escolher_arquivo(pasta)
        if arquivo is None:
            rodada = Rodada(numero, pasta, None, [])
            rodada.motivo = "nenhuma planilha na pasta"
            rodadas.append(rodada)
            log(f"  {pasta.name}: sem planilha.", "aviso")
            continue
        rodada = Rodada(numero, pasta, arquivo, abas_do_arquivo(arquivo))
        rodadas.append(rodada)
        log(f"  Rodada {rodada.numero:>2}: {arquivo.name} — {rodada.situacao}.")
    return rodadas


# ---------------------------------------------------------------------------
#  Leitura do Termo emitido
# ---------------------------------------------------------------------------
#  Rótulos do cabeçalho da aba TERMO, como aparecem nas duas gerações de
#  arquivo. O layout mudou ao longo de 2026 — nos arquivos até a 8ª rodada o
#  cabeçalho está na linha 5, nos de agora na linha 4 — então a posição é
#  **detectada**, nunca assumida. Assumir custou caro: com a linha fixa, a 8ª
#  rodada aparecia com 668 linhas repetidas que não existem.
_ROTULOS_TERMO = {
    "nome": ("NOME",),
    "documento": ("CPF", "CPF/CNPJ", "CNPJ"),
    "identificador": ("N", "N°", "No", "NUMERO", "TITULO", "ID"),
    "vencimento": ("VENCIMENTO",),
    "valor_nominal": ("VALOR NOMINAL",),
    "valor_recompra": ("VALOR RECOMPRA",),
}


def _mapear_cabecalho_termo(tabela):
    """Acha a linha de cabeçalho da aba TERMO e a posição de cada coluna."""
    for numero, linha in tabela.linhas[:20]:
        rotulos = [_sem_acento(c) for c in linha]
        if "NOME" not in rotulos:
            continue
        posicoes = {}
        for campo, aceitos in _ROTULOS_TERMO.items():
            for posicao, rotulo in enumerate(rotulos):
                if rotulo in aceitos or (campo == "documento"
                                         and rotulo.startswith("CPF")):
                    posicoes.setdefault(campo, posicao)
        if "nome" in posicoes and "documento" in posicoes:
            return numero, posicoes
    return None, {}


def ler_termo_emitido(caminho):
    """
    As linhas do Termo que de fato saiu, na ordem em que estão na aba.

    Devolve ``(linhas, identificador_confiavel)``. Cada linha é um dicionário
    com nome, documento, identificador, vencimento e os dois valores — **com as
    repetições**, porque é a linha repetida que revela a cobrança em duplicidade.

    ``identificador_confiavel`` diz se a coluna de número serve como identidade
    do título. Nos arquivos antigos ela **não** serve: o mesmo ``N° 3817908``
    aparece para SARA MARTINS e para LAUDECIR DE ALMEIDA, que são pessoas
    diferentes. Onde ela não serve, a duplicidade é apurada pela assinatura da
    linha (documento, vencimento e valor).
    """
    tabela = leitura.ler_aba(caminho, config.ABA_TERMO)
    cabecalho, posicoes = _mapear_cabecalho_termo(tabela)
    if cabecalho is None:
        return [], False

    def pega(linha, campo):
        posicao = posicoes.get(campo)
        if posicao is None or posicao >= len(linha):
            return None
        return linha[posicao]

    linhas = []
    for numero, bruta in tabela.linhas:
        if numero <= cabecalho:
            continue
        nome = str(pega(bruta, "nome") or "").strip()
        documento = nz.documento(pega(bruta, "documento"))
        if not nome and not documento:
            continue
        identificador = pega(bruta, "identificador")
        linhas.append({
            "linha": numero,
            "identificador": str(identificador).strip() if identificador not in
            (None, "", ".") else "",
            "nome": nome,
            "documento": documento,
            "vencimento": nz.data(pega(bruta, "vencimento")),
            "valor_nominal": nz.valor(pega(bruta, "valor_nominal")),
            "valor_recompra": nz.valor(pega(bruta, "valor_recompra")),
        })

    return linhas, _identificador_confiavel(linhas)


def _identificador_confiavel(linhas):
    """
    A coluna de número identifica o título, ou é outra coisa?

    O teste é direto: se o mesmo número aparece para sacados diferentes, ele não
    é identidade de título. Sem isso, contar repetição por essa coluna produz
    passivo inexistente.
    """
    donos = {}
    for linha in linhas:
        identificador = linha["identificador"]
        if not identificador:
            return False
        dono = linha["documento"] or linha["nome"]
        if donos.setdefault(identificador, dono) != dono:
            return False
    return bool(linhas)


# ---------------------------------------------------------------------------
#  Reprocessamento de uma rodada
# ---------------------------------------------------------------------------
def copiar_para_trabalho(rodada, log=None):
    """Copia o arquivo do OneDrive para a pasta de trabalho (regras 12-A e 15.8)."""
    log = log or (lambda *a, **k: None)
    PASTA_COPIAS.mkdir(parents=True, exist_ok=True)
    destino = PASTA_COPIAS / f"rodada-{rodada.numero:02d}_{rodada.arquivo.name}"
    if not destino.exists() or destino.stat().st_size != rodada.arquivo.stat().st_size:
        log(f"    copiando {rodada.arquivo.name} "
            f"({rodada.arquivo.stat().st_size / 1024 / 1024:.0f} MB)…")
        shutil.copy2(rodada.arquivo, destino)
    rodada.copia = destino
    return destino


def reprocessar(rodada, historico_temporario, log=None):
    """
    Roda o motor corrigido sobre a rodada e compara com o Termo emitido.

    ``historico_temporario`` acumula os títulos das rodadas já processadas, para
    que a trava contra recompra repetida funcione como teria funcionado à época.
    """
    log = log or (lambda *a, **k: None)
    if rodada.situacao == SITUACAO_SEM_ESTRUTURA:
        rodada.motivo = (f"a planilha não tem aba de Termo "
                         f"(abas: {', '.join(rodada.abas) or 'nenhuma'})")
        log(f"  Rodada {rodada.numero}: {rodada.motivo}", "aviso")
        return rodada

    copiar_para_trabalho(rodada, log=log)

    #  O Termo emitido dá para ler nos dois casos, e é ele que revela duplicidade.
    try:
        rodada.emitidos, rodada.id_confiavel = ler_termo_emitido(rodada.copia)
    finally:
        leitura.fechar_cache()

    if not rodada.emitidos:
        rodada.motivo = "não consegui achar o cabeçalho da aba do Termo"
        log(f"  Rodada {rodada.numero}: {rodada.motivo}", "aviso")
        return rodada

    _apurar_duplicidade(rodada)
    log(f"  Rodada {rodada.numero}: Termo emitido com {len(rodada.emitidos)} linhas "
        f"para {rodada.titulos_distintos} títulos"
        + (f" — {len(rodada.duplicados)} linhas repetidas "
           f"(por {rodada.criterio_duplicidade})." if rodada.duplicados else "."),
        "aviso" if rodada.duplicados else "info")

    if rodada.situacao == SITUACAO_SO_TERMO:
        rodada.motivo = ("a planilha não traz as abas VORTX e GRAFENO — só a "
                         "conferência de duplicidade dentro do Termo foi possível")
        _alimentar_historico(historico_temporario, rodada)
        return rodada

    # -- reconciliação completa ---------------------------------------------
    cfg = {
        "template": str(rodada.copia),      # juros e multa da própria rodada
        #  O número vem da pasta, não da aba INFORMAÇÕES: em vários arquivos
        #  antigos essa célula ficou com o valor de outra rodada.
        "numero_rodada": rodada.numero,
        "fonte_vortx": "arquivo", "arquivo_vortx": str(rodada.copia),
        "fonte_grafeno": "arquivo", "arquivo_grafeno": str(rodada.copia),
        "modo_simulacao": True,
        "gravar_snapshot": False,
        "atualizar_historico": False,
        "historico": str(historico_temporario),
    }
    try:
        resultado = motor.conciliar(cfg, log=_log_indentado(log))
    except Exception as erro:                                # noqa: BLE001
        rodada.motivo = f"o motor não conseguiu processar: {erro}"
        log(f"  Rodada {rodada.numero}: {rodada.motivo}", "erro")
        _alimentar_historico(historico_temporario, rodada)
        return rodada

    rodada.parametros = resultado.parametros
    rodada.contagens = dict(resultado.conciliacao.contagens)
    rodada.calculados = {l["id_titulo"]: l for l in resultado.linhas_termo}

    emitidos_unicos = {l["identificador"] for l in rodada.emitidos}
    for identificador, linha in rodada.calculados.items():
        if identificador not in emitidos_unicos:
            rodada.faltantes.append(linha)

    vistos = set()
    for linha in rodada.emitidos:
        identificador = linha["identificador"]
        if identificador in rodada.calculados or identificador in vistos:
            continue
        vistos.add(identificador)
        rodada.indevidos.append(linha)
        rodada.causas[identificador] = _porque_nao(resultado, rodada, identificador)

    log(f"    motor: {len(rodada.calculados)} títulos · "
        f"faltaram {len(rodada.faltantes)} · "
        f"entraram sem dever {len(rodada.indevidos)} · "
        f"repetidos {len(rodada.duplicados)}",
        "aviso" if (rodada.faltantes or rodada.indevidos) else "ok")

    _alimentar_historico(historico_temporario, rodada)
    #  A partir daqui a rodada só carrega texto e números. As duas bases saem
    #  da memória antes de a próxima rodada ser lida.
    del resultado
    return rodada


def _chave_linha(linha):
    """Assinatura de uma linha do Termo, quando o identificador não serve."""
    return (f"{linha['documento']}|{nz.iso(linha['vencimento'])}"
            f"|{nz.valor_texto(linha['valor_nominal'])}")


def _apurar_duplicidade(rodada):
    """
    Linhas cobradas mais de uma vez no mesmo Termo.

    Usa o ``IdTituloVortx`` quando ele é de fato identidade de título; nos
    arquivos antigos, em que a coluna de número se repete entre sacados
    diferentes, cai para a assinatura (documento, vencimento, valor). O critério
    usado vai para o relatório — quem confere precisa saber o que foi contado.
    """
    if rodada.id_confiavel:
        rodada.criterio_duplicidade = "IdTituloVortx"
        chave = lambda l: l["identificador"]              # noqa: E731
    else:
        rodada.criterio_duplicidade = "documento + vencimento + valor"
        chave = _chave_linha

    vistos = {}
    for linha in rodada.emitidos:
        k = chave(linha)
        vistos.setdefault(k, []).append(linha)
    rodada.titulos_distintos = len(vistos)
    rodada.vezes_por_chave = {k: len(v) for k, v in vistos.items()}
    for k, ocorrencias in vistos.items():
        #  A primeira linha é legítima; da segunda em diante é cobrança repetida.
        for linha in ocorrencias[1:]:
            rodada.duplicados.append(dict(linha, chave=k,
                                          vezes=len(ocorrencias)))


def _log_indentado(log):
    def interno(mensagem, nivel="info"):
        if mensagem and mensagem.strip():
            log(f"      {mensagem.strip()}", nivel)
    return interno


def _alimentar_historico(caminho, rodada):
    """
    Registra no histórico temporário o que aquela rodada **de fato** cobrou.

    É o Termo emitido, não o do motor: a rodada seguinte precisa enxergar o que
    o mundo real cobrou, senão a trava contra recompra repetida não reproduz o
    caso da Cíntia.
    """
    from core.historico import Historico
    #  Sem identificador confiável não dá para semear a trava: alimentar o
    #  histórico com números que não são de título barraria a rodada seguinte
    #  por engano. Melhor não travar do que travar o título errado.
    if not rodada.emitidos or not rodada.id_confiavel:
        return
    hist = Historico(caminho)
    parametros = rodada.parametros
    hist.registrar(
        numero=rodada.numero,
        data_recompra=getattr(parametros, "data_recompra", None),
        janela_inicio=getattr(parametros, "janela_inicio", None),
        janela_fim=getattr(parametros, "janela_fim", None),
        titulos=sorted({l["identificador"] for l in rodada.emitidos}),
        valor_total=calculo.somar(l["valor_recompra"] for l in rodada.emitidos),
        arquivo_saida=rodada.arquivo.name if rodada.arquivo else "",
        observacao="Termo emitido, importado no reprocessamento retroativo",
    )


# ---------------------------------------------------------------------------
#  Orquestração
# ---------------------------------------------------------------------------
def rodar(raiz=None, destino=None, apenas=None, log=None):
    """
    Reprocessa todas as rodadas disponíveis e grava a planilha consolidada.

    ``apenas`` limita a números de rodada específicos, útil para conferir uma
    sozinha sem esperar as doze.
    """
    log = log or (lambda *a, **k: None)
    rodadas = varrer(raiz, log=log)
    if apenas:
        alvos = {int(n) for n in apenas}
        rodadas = [r for r in rodadas if r.numero in alvos]

    #  Histórico temporário: o real não pode ser tocado por um reprocessamento.
    with tempfile.TemporaryDirectory(prefix="retroativo_") as temporaria:
        historico_temporario = Path(temporaria) / "historico_retroativo.json"
        log("")
        log("── Reprocessando ─────────────────────────────────────", "destaque")
        for rodada in sorted(rodadas, key=lambda r: r.numero):
            reprocessar(rodada, historico_temporario, log=log)

    destino = Path(destino or config.PASTA_SAIDAS)
    destino.mkdir(parents=True, exist_ok=True)
    caminho = consolidar(rodadas, destino, log=log)
    return rodadas, caminho


# ---------------------------------------------------------------------------
#  Planilha consolidada
# ---------------------------------------------------------------------------
def consolidar(rodadas, pasta_destino, log=None):
    """Uma linha por rodada no resumo, e uma aba de detalhe por tipo de achado."""
    log = log or (lambda *a, **k: None)
    from datetime import date
    destino = Path(pasta_destino) / f"Passivo_Retroativo_{date.today():%Y-%m-%d}.xlsx"

    livro = openpyxl.Workbook()
    livro.remove(livro.active)
    _aba_resumo(livro, rodadas)
    _aba_faltantes(livro, rodadas)
    _aba_indevidos(livro, rodadas)
    _aba_duplicados(livro, rodadas)
    _aba_nao_processadas(livro, rodadas)

    livro.save(destino)
    log("")
    log(f"  Consolidado: {destino.name}", "ok")
    return destino


def _aba_resumo(livro, rodadas):
    aba = _nova_aba(livro, "RESUMO", [
        "Rodada", "Arquivo", "Situação", "Data da recompra", "Janela",
        "Linhas no Termo emitido", "Títulos distintos", "Títulos pelo motor",
        "Faltaram (qtd)", "Faltaram (R$)",
        "Entraram sem dever (qtd)", "Entraram sem dever (R$)",
        "Repetidos (qtd)", "Repetidos (R$)", "Duplicidade apurada por",
    ], larguras=[8, 46, 16, 17, 24, 22, 17, 18, 15, 16, 22, 24, 15, 16, 32])

    for rodada in sorted(rodadas, key=lambda r: r.numero):
        parametros = rodada.parametros
        janela = ""
        if parametros and parametros.janela_inicio and parametros.janela_fim:
            janela = (f"{nz.br(parametros.janela_inicio)} a "
                      f"{nz.br(parametros.janela_fim)}")
        _linha(aba, [
            rodada.numero,
            rodada.arquivo.name if rodada.arquivo else "—",
            rodada.situacao,
            getattr(parametros, "data_recompra", None),
            janela,
            len(rodada.emitidos) or "—",
            rodada.titulos_distintos or "—",
            len(rodada.calculados) if rodada.situacao == SITUACAO_COMPLETA else "—",
            len(rodada.faltantes) if rodada.situacao == SITUACAO_COMPLETA else "—",
            rodada.valor_faltante if rodada.situacao == SITUACAO_COMPLETA else "—",
            len(rodada.indevidos) if rodada.situacao == SITUACAO_COMPLETA else "—",
            rodada.valor_indevido if rodada.situacao == SITUACAO_COMPLETA else "—",
            len(rodada.duplicados),
            rodada.valor_duplicado,
            rodada.criterio_duplicidade or "—",
        ], formatos={4: _DATA})

    #  Totalização só do que foi de fato apurado — somar "—" seria inventar.
    completas = [r for r in rodadas if r.situacao == SITUACAO_COMPLETA]
    com_termo = [r for r in rodadas if r.emitidos]
    aba.append([])
    numero = _linha(aba, [
        "TOTAL", f"{len(completas)} rodada(s) reconciliada(s) de {len(rodadas)}",
        "", None, "", "", "", "",
        sum(len(r.faltantes) for r in completas),
        calculo.somar(r.valor_faltante for r in completas),
        sum(len(r.indevidos) for r in completas),
        calculo.somar(r.valor_indevido for r in completas),
        sum(len(r.duplicados) for r in com_termo),
        calculo.somar(r.valor_duplicado for r in com_termo),
        "",
    ])
    for celula in aba[numero]:
        celula.font = Font(name="Calibri", size=11, bold=True)


def _aba_faltantes(livro, rodadas):
    aba = _nova_aba(livro, "FALTANTES", [
        "Rodada", "IdTituloVortx", "Sacado", "CPF/CNPJ", "Vencimento",
        "Valor nominal", "Valor de recompra", "Status do boleto", "Qtd boletos",
        "Casado na passada",
    ], larguras=[8, 16, 40, 20, 13, 15, 18, 20, 12, 18])
    for rodada in sorted(rodadas, key=lambda r: r.numero):
        for linha in rodada.faltantes:
            _linha(aba, [rodada.numero, linha["id_titulo"], linha["nome"],
                         linha["documento_formatado"], linha["vencimento"],
                         linha["valor_nominal"], linha["valor_recompra"],
                         linha["status"], linha["qtd_boletos"], linha["passada"]],
                   formatos={5: _DATA})
    if aba.max_row == 1:
        _linha(aba, ["— nenhum título deixou de entrar —"])


def _aba_indevidos(livro, rodadas):
    aba = _nova_aba(livro, "ENTRARAM_SEM_DEVER", [
        "Rodada", "IdTituloVortx", "Sacado", "CPF/CNPJ", "Vencimento",
        "Valor cobrado", "Por que o motor não o traria",
    ], larguras=[8, 16, 40, 20, 13, 16, 62])
    for rodada in sorted(rodadas, key=lambda r: r.numero):
        for linha in rodada.indevidos:
            _linha(aba, [rodada.numero, linha["identificador"], linha["nome"],
                         nz.documento_formatado(linha["documento"]),
                         linha["vencimento"], linha["valor_recompra"],
                         rodada.causas.get(linha["identificador"], "—")],
                   formatos={5: _DATA})
    if aba.max_row == 1:
        _linha(aba, ["— nenhum título entrou indevidamente —"])


def _porque_nao(resultado, rodada, identificador):
    """
    A causa de o motor não trazer um título que o Termo trouxe.

    Resolvida na hora, com a base ainda em memória, e guardada como texto — o
    relatório é montado depois que as bases já foram descartadas.
    """
    for casamento, motivo in resultado.selecao.barrados:
        if casamento.titulo.id_titulo == identificador:
            return motivo
    for casamento in resultado.conciliacao.casamentos:
        if casamento.titulo.id_titulo != identificador:
            continue
        if not casamento.casou:
            return f"sem boleto na Grafeno ({casamento.causa})"
        vencimento = casamento.titulo.vencimento
        parametros = rodada.parametros
        if (vencimento and parametros and parametros.janela_inicio
                and not (parametros.janela_inicio <= vencimento <= parametros.janela_fim)):
            return f"vencimento {nz.br(vencimento)} fora da janela da rodada"
        return f"status do boleto é {casamento.status!r}, não 'Aberta (Vencida)'"
    return "o título não existe na base VORTX do próprio arquivo"


def _aba_duplicados(livro, rodadas):
    aba = _nova_aba(livro, "COBRADOS_EM_DUPLICIDADE", [
        "Rodada", "Identificador", "Sacado", "CPF/CNPJ", "Vencimento",
        "Valor da linha repetida", "Vezes no Termo", "Linha na planilha",
        "Apurado por",
    ], larguras=[8, 16, 40, 20, 13, 24, 15, 17, 32])
    for rodada in sorted(rodadas, key=lambda r: r.numero):
        for linha in rodada.duplicados:
            _linha(aba, [rodada.numero, linha["identificador"], linha["nome"],
                         nz.documento_formatado(linha["documento"]),
                         linha["vencimento"], linha["valor_recompra"],
                         linha["vezes"], linha["linha"],
                         rodada.criterio_duplicidade],
                   formatos={5: _DATA})
    if aba.max_row == 1:
        _linha(aba, ["— nenhuma linha repetida em nenhum Termo —"])


def _aba_nao_processadas(livro, rodadas):
    aba = _nova_aba(livro, "NAO_RECONCILIADAS", [
        "Rodada", "Pasta", "Arquivo", "Abas encontradas", "Por que não",
    ], larguras=[8, 26, 46, 50, 60])
    for rodada in sorted(rodadas, key=lambda r: r.numero):
        if rodada.situacao == SITUACAO_COMPLETA and not rodada.motivo:
            continue
        _linha(aba, [rodada.numero, rodada.pasta.name,
                     rodada.arquivo.name if rodada.arquivo else "—",
                     ", ".join(rodada.abas) or "—",
                     rodada.motivo or rodada.situacao])
    if aba.max_row == 1:
        _linha(aba, ["— todas as rodadas foram reconciliadas —"])
    for celula in aba[1]:
        celula.alignment = Alignment(horizontal="center", wrap_text=True)
