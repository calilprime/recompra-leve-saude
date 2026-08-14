# -*- coding: utf-8 -*-
"""
Leitura de tabela em .xlsx ou .csv.
===================================

Camada baixa, compartilhada por :mod:`core.fonte_vortx` e
:mod:`core.fonte_grafeno`. Devolve as linhas como tuplas, junto com o cabeçalho
— guardar dicionário para 142 mil boletos custa memória à toa, e as duas abas
precisam ser regravadas coluna a coluna, na ordem original.

Regra herdada do defeito 2.1: **nunca** parar a leitura num número de linha
fixo. A varredura vai até o fim do que existe e a última linha com dado é
descoberta, não arbitrada.
"""

import csv
import sys
from pathlib import Path

import openpyxl

CODIFICACOES = ("utf-8-sig", "utf-8", "cp1252", "latin-1")


class Tabela:
    """Cabeçalho + linhas de uma planilha, com acesso por nome de coluna."""

    __slots__ = ("cabecalho", "linhas", "indice", "origem", "aba")

    def __init__(self, cabecalho, linhas, origem="", aba=""):
        self.cabecalho = list(cabecalho)
        self.linhas = linhas
        self.origem = str(origem)
        self.aba = aba
        self.indice = {}
        for posicao, nome in enumerate(self.cabecalho):
            if nome not in self.indice:          # 1ª ocorrência vence
                self.indice[nome] = posicao

    def __len__(self):
        return len(self.linhas)

    def posicao(self, coluna, obrigatoria=True):
        """Índice da coluna pelo nome do cabeçalho."""
        if coluna in self.indice:
            return self.indice[coluna]
        #  Tolera diferença de espaço e de caixa ("VENCIMENTO " no Termo).
        alvo = str(coluna).strip().casefold()
        for nome, posicao in self.indice.items():
            if str(nome).strip().casefold() == alvo:
                return posicao
        if obrigatoria:
            raise KeyError(
                f"Coluna {coluna!r} não existe em {self.origem} / {self.aba}. "
                f"Colunas encontradas: {', '.join(map(str, self.cabecalho[:60]))}"
            )
        return None

    def valor(self, linha, coluna, padrao=None):
        """Valor da coluna nomeada, numa linha já lida."""
        posicao = self.posicao(coluna, obrigatoria=False)
        if posicao is None or posicao >= len(linha):
            return padrao
        v = linha[posicao]
        return padrao if v is None else v


# ---------------------------------------------------------------------------
#  Excel
# ---------------------------------------------------------------------------
#  Abrir um .xlsx de 50 MB custa alguns segundos e o descompactar da aba custa
#  bem mais. Quando as duas abas vêm do mesmo arquivo — que é o caso ao
#  reprocessar uma rodada antiga — abrir uma vez só corta o tempo pela metade.
_ABERTOS = {}


def _abrir(caminho, data_only):
    chave = (str(caminho), Path(caminho).stat().st_mtime_ns, bool(data_only))
    if chave not in _ABERTOS:
        fechar_cache()
        _ABERTOS[chave] = openpyxl.load_workbook(
            caminho, read_only=True, data_only=data_only)
    return _ABERTOS[chave]


def fechar_cache():
    """Fecha os arquivos abertos. Chamar ao fim de cada rodada."""
    for wb in _ABERTOS.values():
        try:
            wb.close()
        except Exception:
            pass
    _ABERTOS.clear()


def ler_aba(caminho, aba, colunas_sentinela=(), data_only=True, log=None):
    """
    Lê uma aba inteira de um .xlsx.

    ``colunas_sentinela`` são as colunas que definem "esta linha tem dado".
    A varredura percorre a aba toda e corta o rabo de linhas em branco que as
    fórmulas estendidas deixam para trás — foi o que escondeu as sete linhas
    órfãs da 12ª rodada.
    """
    caminho = Path(caminho)
    wb = _abrir(caminho, data_only)
    aba = _nome_real(wb, aba, caminho)
    ws = wb[aba]
    iterador = ws.iter_rows(values_only=True)
    try:
        cabecalho = [c for c in next(iterador)]
    except StopIteration:
        return Tabela([], [], caminho, aba)

    while cabecalho and cabecalho[-1] is None:
        cabecalho.pop()

    tabela = Tabela(cabecalho, [], caminho, aba)
    sentinelas = [tabela.posicao(c) for c in colunas_sentinela] or None

    #  A linha de dado não é cortada na largura do cabeçalho: em várias abas o
    #  cabeçalho é mais estreito que os dados (o Termo tem só B1 preenchido) e
    #  cortar esconderia coluna inteira.
    linhas = []
    for numero, linha in enumerate(iterador, start=2):
        if sentinelas is None:
            if all(v is None or v == "" for v in linha):
                continue
        else:
            if all(
                p >= len(linha) or linha[p] is None or linha[p] == ""
                for p in sentinelas
            ):
                continue
        linhas.append((numero, linha))
    tabela.linhas = linhas
    if log:
        log(f"    {aba}: {len(linhas)} linhas com dado "
            f"(varridas até a linha {ws.max_row}).")
    return tabela


def _nome_real(wb, aba, caminho):
    """A acentuação do nome da aba pode chegar diferente; compara sem ela."""
    if aba in wb.sheetnames:
        return aba
    from core.normalizacao import sem_acento
    alvo = sem_acento(aba)
    achou = [n for n in wb.sheetnames if sem_acento(n) == alvo]
    if not achou:
        raise KeyError(f"A aba {aba!r} não existe em {Path(caminho).name}. "
                       f"Abas: {', '.join(wb.sheetnames)}")
    return achou[0]


def ler_celulas(caminho, aba, celulas, data_only=True):
    """Lê células soltas de uma aba (usado na aba INFORMAÇÕES)."""
    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=data_only)
    try:
        ws = wb[_nome_real(wb, aba, caminho)]
        #  read_only não aceita ws['D8'] em qualquer ordem: lê o bloco todo.
        maximo = max(_linha_de(c) for c in celulas)
        grade = {}
        for numero, linha in enumerate(
            ws.iter_rows(min_row=1, max_row=maximo, values_only=True), start=1
        ):
            grade[numero] = linha
        resultado = {}
        for c in celulas:
            coluna, numero = _coluna_de(c), _linha_de(c)
            linha = grade.get(numero) or ()
            resultado[c] = linha[coluna - 1] if coluna - 1 < len(linha) else None
        return resultado
    finally:
        wb.close()


def _coluna_de(celula):
    letras = "".join(ch for ch in celula if ch.isalpha()).upper()
    indice = 0
    for ch in letras:
        indice = indice * 26 + (ord(ch) - 64)
    return indice


def _linha_de(celula):
    return int("".join(ch for ch in celula if ch.isdigit()))


# ---------------------------------------------------------------------------
#  CSV
# ---------------------------------------------------------------------------
def ler_csv(caminho, colunas_sentinela=(), log=None):
    """Lê um .csv detectando codificação e separador."""
    caminho = Path(caminho)
    bruto = caminho.read_bytes()
    conteudo = None
    for codificacao in CODIFICACOES:
        try:
            conteudo = bruto.decode(codificacao)
            break
        except UnicodeDecodeError:
            continue
    if conteudo is None:
        conteudo = bruto.decode("latin-1", errors="replace")

    amostra = conteudo[:8192]
    try:
        separador = csv.Sniffer().sniff(amostra, delimiters=";,\t|").delimiter
    except csv.Error:
        separador = ";" if amostra.count(";") >= amostra.count(",") else ","

    limite_antigo = csv.field_size_limit()
    csv.field_size_limit(min(sys.maxsize, 2 ** 31 - 1))
    try:
        leitor = csv.reader(conteudo.splitlines(), delimiter=separador)
        try:
            cabecalho = next(leitor)
        except StopIteration:
            return Tabela([], [], caminho, "csv")
        tabela = Tabela(cabecalho, [], caminho, "csv")
        sentinelas = [tabela.posicao(c) for c in colunas_sentinela] or None

        linhas = []
        for numero, linha in enumerate(leitor, start=2):
            if sentinelas is None:
                if all(v == "" for v in linha):
                    continue
            else:
                if all(p >= len(linha) or linha[p] == "" for p in sentinelas):
                    continue
            linhas.append((numero, tuple(linha)))
        tabela.linhas = linhas
    finally:
        csv.field_size_limit(limite_antigo)

    if log:
        log(f"    {caminho.name}: {len(tabela.linhas)} linhas "
            f"(separador {separador!r}).")
    return tabela


def ler(caminho, aba=None, colunas_sentinela=(), log=None):
    """Despacha entre .xlsx e .csv pelo sufixo do arquivo."""
    caminho = Path(caminho)
    if not caminho.exists():
        raise FileNotFoundError(f"Arquivo não encontrado: {caminho}")
    if caminho.suffix.lower() in (".csv", ".txt"):
        return ler_csv(caminho, colunas_sentinela, log=log)
    if aba is None:
        raise ValueError("Para arquivo Excel é preciso informar a aba.")
    return ler_aba(caminho, aba, colunas_sentinela, log=log)
